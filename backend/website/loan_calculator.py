import math
import io
from datetime import datetime, date
from dateutil.relativedelta import relativedelta
from flask import Blueprint, request, jsonify, send_file
from flask_jwt_extended import jwt_required

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, numbers
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.units import mm
from reportlab.lib.enums import TA_LEFT, TA_RIGHT, TA_CENTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from website.models import Application, Company

loan_calc_bp = Blueprint('loan_calculator', __name__, url_prefix='/api')


class LoanCalculator:
    """Core mathematical engine for loan disbursements, interest, amortization, and penalties."""

    @staticmethod
    def calculate_breakdown(
        principal: float,
        interest_rate: float = 10.0,
        duration_months: int = 1,
        processing_fee_rate: float = 3.0,
        access_fee_rate: float = 0.0,
        legal_fee: float = 0.0,
        rate_type: str = "monthly",  # "monthly" or "annual"
        method: str = "flat"         # "flat" or "reducing"
    ) -> dict:
        """
        Compute full financial breakdown for a loan application.
        """
        principal = float(principal)
        interest_rate = float(interest_rate)
        duration_months = max(1, int(duration_months))
        processing_fee_rate = float(processing_fee_rate)
        access_fee_rate = float(access_fee_rate)
        legal_fee = float(legal_fee)

        # 1. Deductions / Upfront Fees
        processing_fee = round(principal * (processing_fee_rate / 100.0), 2)
        access_fee = round(principal * (access_fee_rate / 100.0), 2)
        total_deductions = round(processing_fee + access_fee + legal_fee, 2)
        net_disbursement = round(max(0.0, principal - total_deductions), 2)

        # 2. Interest Calculation
        if rate_type == "annual":
            monthly_rate = (interest_rate / 100.0) / 12.0
        else:
            monthly_rate = interest_rate / 100.0

        if method == "reducing" and duration_months > 1 and monthly_rate > 0:
            # Equated Monthly Installment (EMI) formula: P * r * (1 + r)^n / ((1 + r)^n - 1)
            emi = principal * (monthly_rate * math.pow(1 + monthly_rate, duration_months)) / (math.pow(1 + monthly_rate, duration_months) - 1)
            monthly_installment = round(emi, 2)
            total_repayment = round(monthly_installment * duration_months, 2)
            total_interest = round(total_repayment - principal, 2)
        else:
            # Flat Rate calculation
            total_interest = round(principal * monthly_rate * duration_months, 2)
            total_repayment = round(principal + total_interest, 2)
            monthly_installment = round(total_repayment / duration_months, 2)

        return {
            "principal": principal,
            "interest_rate": interest_rate,
            "rate_type": rate_type,
            "method": method,
            "duration_months": duration_months,
            "processing_fee": processing_fee,
            "processing_fee_rate": processing_fee_rate,
            "access_fee": access_fee,
            "access_fee_rate": access_fee_rate,
            "legal_fee": legal_fee,
            "total_deductions": total_deductions,
            "net_disbursement": net_disbursement,
            "total_interest": total_interest,
            "total_repayment": total_repayment,
            "monthly_installment": monthly_installment
        }

    @staticmethod
    def generate_amortization_schedule(
        principal: float,
        interest_rate: float = 10.0,
        duration_months: int = 1,
        start_date: date = None,
        rate_type: str = "monthly",
        method: str = "flat"
    ) -> list:
        """
        Generate month-by-month amortization schedule breakdown.
        """
        if start_date is None:
            start_date = date.today()

        principal = float(principal)
        interest_rate = float(interest_rate)
        duration_months = max(1, int(duration_months))

        if rate_type == "annual":
            monthly_rate = (interest_rate / 100.0) / 12.0
        else:
            monthly_rate = interest_rate / 100.0

        schedule = []
        current_balance = principal

        if method == "reducing" and duration_months > 1 and monthly_rate > 0:
            emi = principal * (monthly_rate * math.pow(1 + monthly_rate, duration_months)) / (math.pow(1 + monthly_rate, duration_months) - 1)
            installment_amount = round(emi, 2)

            for i in range(1, duration_months + 1):
                due_date = start_date + relativedelta(months=+i)
                interest_payment = round(current_balance * monthly_rate, 2)
                
                if i == duration_months:
                    # Final installment adjustment to clear exact remaining balance
                    principal_payment = round(current_balance, 2)
                    installment_amount = round(principal_payment + interest_payment, 2)
                    closing_balance = 0.0
                else:
                    principal_payment = round(installment_amount - interest_payment, 2)
                    closing_balance = round(max(0.0, current_balance - principal_payment), 2)

                schedule.append({
                    "installment_number": i,
                    "due_date": due_date.strftime("%d %B %Y"),
                    "due_date_iso": due_date.isoformat(),
                    "opening_balance": current_balance,
                    "principal_payment": principal_payment,
                    "interest_payment": interest_payment,
                    "total_payment": installment_amount,
                    "closing_balance": closing_balance
                })
                current_balance = closing_balance
        else:
            # Flat Rate schedule
            total_interest = round(principal * monthly_rate * duration_months, 2)
            principal_per_month = round(principal / duration_months, 2)
            interest_per_month = round(total_interest / duration_months, 2)
            installment_amount = round(principal_per_month + interest_per_month, 2)

            for i in range(1, duration_months + 1):
                due_date = start_date + relativedelta(months=+i)
                
                if i == duration_months:
                    principal_payment = round(current_balance, 2)
                    installment_amount = round(principal_payment + interest_per_month, 2)
                    closing_balance = 0.0
                else:
                    principal_payment = principal_per_month
                    closing_balance = round(max(0.0, current_balance - principal_payment), 2)

                schedule.append({
                    "installment_number": i,
                    "due_date": due_date.strftime("%d %B %Y"),
                    "due_date_iso": due_date.isoformat(),
                    "opening_balance": current_balance,
                    "principal_payment": principal_payment,
                    "interest_payment": interest_per_month,
                    "total_payment": installment_amount,
                    "closing_balance": closing_balance
                })
                current_balance = closing_balance

        return schedule

    @staticmethod
    def calculate_penalties(
        overdue_amount: float,
        due_date: datetime,
        as_of_date: datetime = None,
        grace_period_days: int = 3,
        daily_penalty_rate: float = 0.005, # 0.5% daily
        fixed_collection_fee: float = 0.0
    ) -> dict:
        """
        Calculate late payment fees and accrued penalties for overdue accounts.
        """
        if as_of_date is None:
            as_of_date = datetime.utcnow()

        overdue_amount = float(overdue_amount)
        if overdue_amount <= 0:
            return {
                "is_overdue": False,
                "days_overdue": 0,
                "penalty_fee": 0.0,
                "collection_fee": 0.0,
                "total_overdue_payable": 0.0
            }

        if isinstance(due_date, str):
            try:
                due_date = datetime.fromisoformat(due_date.replace('Z', '+00:00'))
            except Exception:
                due_date = datetime.strptime(due_date, "%Y-%m-%d")

        days_overdue = (as_of_date.date() - due_date.date()).days

        if days_overdue <= grace_period_days:
            return {
                "is_overdue": False,
                "days_overdue": max(0, days_overdue),
                "grace_period_remaining": max(0, grace_period_days - days_overdue),
                "penalty_fee": 0.0,
                "collection_fee": 0.0,
                "total_overdue_payable": overdue_amount
            }

        billable_days = days_overdue - grace_period_days
        penalty_fee = round(overdue_amount * (daily_penalty_rate * billable_days), 2)
        total_overdue_payable = round(overdue_amount + penalty_fee + fixed_collection_fee, 2)

        return {
            "is_overdue": True,
            "days_overdue": days_overdue,
            "billable_penalty_days": billable_days,
            "daily_rate_percent": daily_penalty_rate * 100,
            "penalty_fee": penalty_fee,
            "collection_fee": fixed_collection_fee,
            "total_overdue_payable": total_overdue_payable
        }


@loan_calc_bp.route('/loans/calculate', methods=['POST', 'GET'])
def simulate_loan():
    """
    Public/Protected calculation API for frontend loan calculators and previews.
    Supports query parameters or JSON body.
    """
    data = request.get_json(silent=True) or {}
    if not data and request.args:
        data = request.args.to_dict()

    try:
        principal = float(data.get('amount') or data.get('principal') or 10000.0)
        interest_rate = float(data.get('interest_rate') or 10.0)
        duration_months = int(data.get('duration_months') or data.get('duration') or 1)
        processing_fee_rate = float(data.get('processing_fee_rate') or 3.0)
        access_fee_rate = float(data.get('access_fee_rate') or 0.0)
        legal_fee = float(data.get('legal_fee') or 0.0)
        rate_type = data.get('rate_type', 'monthly')
        method = data.get('method', 'flat')

        # Company override if company code provided
        company_code = data.get('company')
        if company_code:
            company = Company.query.filter_by(code=company_code.lower()).first()
            if company:
                if 'interest_rate' not in data:
                    interest_rate = company.interest_rate
                if 'processing_fee_rate' not in data:
                    processing_fee_rate = company.processing_fee_rate
                if 'access_fee_rate' not in data:
                    access_fee_rate = company.access_fee_rate
                if 'legal_fee' not in data:
                    legal_fee = company.legal_fee

        breakdown = LoanCalculator.calculate_breakdown(
            principal=principal,
            interest_rate=interest_rate,
            duration_months=duration_months,
            processing_fee_rate=processing_fee_rate,
            access_fee_rate=access_fee_rate,
            legal_fee=legal_fee,
            rate_type=rate_type,
            method=method
        )

        schedule = LoanCalculator.generate_amortization_schedule(
            principal=principal,
            interest_rate=interest_rate,
            duration_months=duration_months,
            rate_type=rate_type,
            method=method
        )

        return jsonify({
            "success": True,
            "breakdown": breakdown,
            "schedule": schedule
        }), 200

    except Exception as e:
        return jsonify({"message": "Invalid loan parameters", "error": str(e)}), 400


def _get_loan_data_from_request():
    data = request.get_json(silent=True) or {}
    if not data and request.args:
        data = request.args.to_dict()

    principal = float(data.get('amount') or data.get('principal') or 10000.0)
    interest_rate = float(data.get('interest_rate') or 10.0)
    duration_months = int(data.get('duration_months') or data.get('duration') or 1)
    processing_fee_rate = float(data.get('processing_fee_rate') or 3.0)
    access_fee_rate = float(data.get('access_fee_rate') or 0.0)
    legal_fee = float(data.get('legal_fee') or 0.0)
    rate_type = data.get('rate_type', 'monthly')
    method = data.get('method', 'flat')
    
    company_code = data.get('company')
    if company_code:
        company = Company.query.filter_by(code=company_code.lower()).first()
        if company:
            if 'interest_rate' not in data: interest_rate = company.interest_rate
            if 'processing_fee_rate' not in data: processing_fee_rate = company.processing_fee_rate
            if 'access_fee_rate' not in data: access_fee_rate = company.access_fee_rate
            if 'legal_fee' not in data: legal_fee = company.legal_fee

    breakdown = LoanCalculator.calculate_breakdown(
        principal=principal, interest_rate=interest_rate, duration_months=duration_months,
        processing_fee_rate=processing_fee_rate, access_fee_rate=access_fee_rate,
        legal_fee=legal_fee, rate_type=rate_type, method=method
    )

    schedule = LoanCalculator.generate_amortization_schedule(
        principal=principal, interest_rate=interest_rate, duration_months=duration_months,
        rate_type=rate_type, method=method
    )
    return breakdown, schedule


@loan_calc_bp.route('/loans/amortization/pdf', methods=['POST'])
def export_amortization_pdf():
    try:
        breakdown, schedule = _get_loan_data_from_request()
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, 
            pagesize=A4,
            leftMargin=20 * mm,
            rightMargin=20 * mm,
            topMargin=15 * mm,
            bottomMargin=15 * mm
        )
        styles = getSampleStyleSheet()
        elements = []
        
        # Custom Palette
        c_primary = colors.HexColor("#0F2C59")      # Deep Navy
        c_secondary = colors.HexColor("#D97706")    # Rich Gold
        c_dark = colors.HexColor("#1E293B")         # Charcoal Dark
        c_light = colors.HexColor("#F8FAFC")        # Soft Ice
        c_border = colors.HexColor("#E2E8F0")       # Border Gray
        c_muted = colors.HexColor("#64748B")        # Slate Muted
        
        title_style = ParagraphStyle(
            'CompanyTitle', fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=c_primary, alignment=TA_LEFT
        )
        subtitle_style = ParagraphStyle(
            'SubTitle', fontName='Helvetica', fontSize=8, leading=11, textColor=c_muted, alignment=TA_LEFT
        )
        doc_header_style = ParagraphStyle(
            'DocHeader', fontName='Helvetica-Bold', fontSize=13, leading=16, textColor=c_secondary, alignment=TA_RIGHT
        )
        
        # HEADER
        header_data = [
            [
                [Paragraph("LeoCap Investments", title_style), Paragraph("Westlands, Nairobi, Kenya<br/>hello@leocapinvest.co.ke | +254 700 000 000", subtitle_style)],
                [Paragraph("AMORTIZATION SCHEDULE", doc_header_style), Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", ParagraphStyle('meta', fontName='Helvetica', fontSize=8, textColor=c_muted, alignment=TA_RIGHT))]
            ]
        ]
        
        t_header = Table(header_data, colWidths=['60%', '40%'])
        t_header.setStyle(TableStyle([
            ('ALIGN', (0,0), (0,0), 'LEFT'),
            ('ALIGN', (1,0), (1,0), 'RIGHT'),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 10),
            ('LINEBELOW', (0,0), (-1,-1), 1, c_primary),
        ]))
        elements.append(t_header)
        elements.append(Spacer(1, 15))
        
        # SUMMARY BLOCK
        summary_data = [
            ["Principal Amount:", f"{breakdown['principal']:,.2f} KES", "Interest Rate:", f"{breakdown['interest_rate']}% ({breakdown['rate_type']})"],
            ["Duration:", f"{breakdown['duration_months']} Months", "Calculation Method:", f"{breakdown['method'].title()}"],
            ["Total Deductions:", f"{breakdown['total_deductions']:,.2f} KES", "Net Disbursement:", f"{breakdown['net_disbursement']:,.2f} KES"]
        ]
        
        t_summary = Table(summary_data, colWidths=['20%', '30%', '20%', '30%'])
        t_summary.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), c_light),
            ('TEXTCOLOR', (0,0), (0,-1), c_muted),
            ('TEXTCOLOR', (2,0), (2,-1), c_muted),
            ('TEXTCOLOR', (1,0), (1,-1), c_dark),
            ('TEXTCOLOR', (3,0), (3,-1), c_dark),
            ('FONTNAME', (0,0), (-1,-1), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,-1), 9),
            ('PADDING', (0,0), (-1,-1), 8),
            ('BOX', (0,0), (-1,-1), 1, c_border),
            ('INNERGRID', (0,0), (-1,-1), 0.5, c_border),
        ]))
        elements.append(t_summary)
        elements.append(Spacer(1, 20))
        
        # SCHEDULE TABLE
        th_style = ParagraphStyle('THStyle', fontName='Helvetica-Bold', fontSize=8, textColor=colors.white, alignment=TA_CENTER)
        
        table_data = [[
            Paragraph("MONTH", th_style), 
            Paragraph("DUE DATE", th_style), 
            Paragraph("OPENING BAL", th_style), 
            Paragraph("PRINCIPAL", th_style), 
            Paragraph("INTEREST", th_style), 
            Paragraph("PAYMENT", th_style), 
            Paragraph("CLOSING BAL", th_style)
        ]]
        
        for idx, row in enumerate(schedule):
            bg = colors.white if idx % 2 == 0 else c_light
            table_data.append([
                str(row['installment_number']),
                row['due_date'],
                f"{row['opening_balance']:,.2f}",
                f"{row['principal_payment']:,.2f}",
                f"{row['interest_payment']:,.2f}",
                f"{row['total_payment']:,.2f}",
                f"{row['closing_balance']:,.2f}"
            ])
            
        t_schedule = Table(table_data, colWidths=['10%', '15%', '15%', '15%', '15%', '15%', '15%'])
        t_schedule.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), c_primary),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,1), (-1,-1), 8),
            ('TEXTCOLOR', (0,1), (-1,-1), c_dark),
            ('GRID', (0,0), (-1,-1), 0.5, c_border),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, c_light]),
            ('PADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(t_schedule)
        
        doc.build(elements)
        buffer.seek(0)
        
        return send_file(buffer, as_attachment=True, download_name='amortization_schedule.pdf', mimetype='application/pdf')
    except Exception as e:
        return jsonify({"message": "Failed to generate PDF", "error": str(e)}), 400


@loan_calc_bp.route('/loans/amortization/excel', methods=['POST'])
def export_amortization_excel():
    try:
        breakdown, schedule = _get_loan_data_from_request()
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Amortization Schedule"
        
        headers = ['Month', 'Due Date', 'Opening Balance', 'Principal', 'Interest', 'Total Payment', 'Closing Balance']
        ws.append(headers)
        
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4F81BD")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            
        for row_data in schedule:
            ws.append([
                row_data['installment_number'],
                row_data['due_date'],
                row_data['opening_balance'],
                row_data['principal_payment'],
                row_data['interest_payment'],
                row_data['total_payment'],
                row_data['closing_balance']
            ])
            
        # Format currency and auto-size
        for row in ws.iter_rows(min_row=2, max_row=len(schedule)+1, min_col=3, max_col=7):
            for cell in row:
                cell.number_format = '#,##0.00'
                
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
            
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        return send_file(buffer, as_attachment=True, download_name='amortization_schedule.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e:
        return jsonify({"message": "Failed to generate Excel", "error": str(e)}), 400


@loan_calc_bp.route('/loans/calculate-gross', methods=['POST', 'GET'])
def calculate_gross_loan():
    """Calculate the required Gross Loan amount to receive a specific Net Disbursement."""
    data = request.get_json(silent=True) or {}
    if not data and request.args:
        data = request.args.to_dict()

    try:
        net_amount = float(data.get('net_amount') or 0.0)
        loan_type = data.get('loan_type', 'personal') # 'personal' or 'salary_advance'
        
        # Defaults based on user prompt logic
        if loan_type == 'salary_advance':
            # Legal = 500, Processing = 3%, Access = 2% -> Total Fees = 5%
            gross_amount = (net_amount + 500) / 0.95
        else:
            # Personal: Legal = 1000, Processing = 3%, Access = 0% -> Total Fees = 3%
            gross_amount = (net_amount + 1000) / 0.97
            
        return jsonify({
            "success": True, 
            "data": {
                "net_amount": net_amount,
                "gross_amount": round(gross_amount, 2),
                "loan_type": loan_type
            }
        }), 200
        
    except Exception as e:
        return jsonify({"message": "Failed to calculate gross loan", "error": str(e)}), 400

@loan_calc_bp.route('/loans/simulate-salary-advance', methods=['POST', 'GET'])
def simulate_salary_advance():
    """Simulate Salary Advance Loan based on 3 months net salary."""
    data = request.get_json(silent=True) or {}
    if not data and request.args:
        data = request.args.to_dict()

    try:
        m1 = float(data.get('month1', 0.0))
        m2 = float(data.get('month2', 0.0))
        m3 = float(data.get('month3', 0.0))
        
        avg_salary = (m1 + m2 + m3) / 3.0
        
        # Max repayment is 1/3 of average salary
        max_repayment = avg_salary / 3.0
        
        # Gross Loan Amount = Repayment / 1.1 (Assuming 10% rate)
        gross_loan = max_repayment / 1.1
        
        # Fees
        processing = gross_loan * 0.03
        access = gross_loan * 0.02
        legal = 500.0
        
        disbursement = gross_loan - processing - access - legal
        interest = gross_loan * 0.10
        
        return jsonify({
            "success": True, 
            "data": {
                "average_salary": round(avg_salary, 2),
                "loan_amount": round(gross_loan, 2),
                "processing_fees": round(processing, 2),
                "access_fees": round(access, 2),
                "legal_charge": round(legal, 2),
                "disbursement_amount": round(disbursement, 2),
                "interest": round(interest, 2),
                "repayment_amount": round(max_repayment, 2)
            }
        }), 200
        
    except Exception as e:
        return jsonify({"message": "Failed to simulate salary advance", "error": str(e)}), 400

@loan_calc_bp.route('/applications/<application_id>/amortization', methods=['GET'])
@jwt_required()
def get_application_amortization(application_id):
    """Retrieve amortization schedule for an existing loan application."""
    app = Application.query.get(application_id)
    if not app:
        return jsonify({"message": "Application not found"}), 404

    try:
        principal = float(app.amount_applied or 0)
        interest_rate = float(app.interest_rate or 10)
        start_date = app.created_at.date() if app.created_at else date.today()

        schedule = LoanCalculator.generate_amortization_schedule(
            principal=principal,
            interest_rate=interest_rate,
            duration_months=1,
            start_date=start_date
        )

        return jsonify({
            "success": True,
            "application_id": app.id,
            "loan_type": app.loan_type,
            "repayment_status": app.repayment_status,
            "amount_applied": principal,
            "repayment_amount": float(app.repayment_amount or 0),
            "due_date": app.due_date.isoformat() if app.due_date else None,
            "schedule": schedule
        }), 200

    except Exception as e:
        return jsonify({"message": "Failed to calculate amortization", "error": str(e)}), 500
